"""
pubstyle.report — assemble a Results section (.docx) with numbered three-line tables and figures.

    doc = ResultsDoc("en")                     # or "zh"
    doc.title("3  Results")                    # zh: "3  结果"
    doc.notice()                               # simulated-data disclaimer (use whenever data are simulated)
    doc.heading("3.1  Participants", 2)
    doc.para("Fifty participants ...")
    doc.figure("figs/fig1.png", "Participant flow. ...")      # -> "Figure 1. Participant flow. ..."
    doc.table(df, "Baseline characteristics", note="Mean ± SD or n (%).")   # -> "Table 1. ..."
    doc.save("results.docx")

Typography (fixed by language):
    zh  body 宋体 + Times New Roman 10.5 pt; headings 黑体; captions "图 1  ..." / "表 1  ..."; table font 7–8 pt
    en  Times New Roman throughout; headings bold; captions "Figure 1. ..." / "Table 1. ..." (label in bold)
Three-line tables: thick top and bottom rules, thin rule under the header, no vertical lines, no shading.

Chinese degree thesis (学位论文) layout: ResultsDoc("zh", profile="thesis", chapter=4)
    body 小四 (12 pt) 宋体/Times New Roman, 1.5 line spacing, 2.5 cm margins, tables 五号 (10.5 pt),
    figures and tables numbered by chapter: "图 4-1" / "表 4-1" (en: "Figure 4.1" / "Table 4.1").
"""
from __future__ import annotations

from pathlib import Path
from typing import List, Optional, Sequence

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .style import t


class ResultsDoc:
    def __init__(self, lang: str = "zh", profile: str = "journal", chapter: Optional[int] = None):
        if lang not in ("zh", "en"):
            raise ValueError("lang must be 'zh' or 'en'")
        if profile not in ("journal", "thesis"):
            raise ValueError("profile must be 'journal' or 'thesis'")
        self.lang = lang
        self.profile = profile
        self.chapter = chapter
        thesis = profile == "thesis"
        self.body_pt = 12.0 if thesis else 10.5
        self.line_spacing = 1.5 if thesis else 1.35
        self.table_pt = 10.5 if thesis else 8
        self.doc = Document()
        s = self.doc.sections[0]
        s.page_width, s.page_height = Cm(21), Cm(29.7)
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(2.5 if thesis else 2.0)
        st = self.doc.styles["Normal"]
        st.font.name = "Times New Roman"
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        self.n_fig = 0
        self.n_tab = 0

    # ------------------------------------------------------------------ fonts
    def _font(self, run, size=10.5, bold=False, italic=False, color=None, heading=False):
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.name = "Times New Roman"
        east = "黑体" if (heading and self.lang == "zh") else "宋体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), east)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)

    def _p(self, text="", size=None, bold=False, italic=False, color=None, heading=False, after=4, indent=False, align=None, keep_next=False):
        p = self.doc.add_paragraph()
        if text:
            self._font(p.add_run(text), size or self.body_pt, bold, italic, color, heading)
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.line_spacing = self.line_spacing
        if indent and self.lang == "zh":
            p.paragraph_format.first_line_indent = Cm(0.74)
        if align == "center":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = keep_next
        return p

    # ------------------------------------------------------------------ text blocks
    def title(self, text: str):
        return self._p(text, 16, True, heading=True, after=4)

    def heading(self, text: str, level: int = 2):
        size = {1: 14, 2: 12, 3: 11}[level]
        p = self._p(text, size, True, heading=True, after=6, keep_next=True)
        p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
        return p

    def para(self, text: str, indent: bool = True):
        return self._p(text, indent=indent)

    def notice(self, text: Optional[str] = None):
        return self._p(text or t("disclaimer", self.lang), 9, italic=True, color="9A3412", after=8)

    # ------------------------------------------------------------------ captions
    def num_label(self, n: int) -> str:
        """Displayed number: '3' (journal) or '4-3' / '4.3' (thesis, numbered by chapter)."""
        if self.chapter is None:
            return str(n)
        return f"{self.chapter}-{n}" if self.lang == "zh" else f"{self.chapter}.{n}"

    def _caption(self, kind: str, n: int, body: str, above: bool):
        lab = t(kind, self.lang)
        num = self.num_label(n)
        p = self.doc.add_paragraph()
        size = 10.5 if self.profile == "thesis" else (9.5 if above else 8.5)
        if self.lang == "zh":
            self._font(p.add_run(f"{lab} {num}  "), size, True, heading=True)
            self._font(p.add_run(body), size, above and self.profile != "thesis")
        else:
            self._font(p.add_run(f"{lab} {num}. "), size, True)
            self._font(p.add_run(body), size, False)
        p.paragraph_format.space_after = Pt(3 if above else 10)
        p.paragraph_format.keep_with_next = above
        return p

    # ------------------------------------------------------------------ figure
    def figure(self, png_path: str | Path, caption: str, width_cm: float = 17.0) -> int:
        self.n_fig += 1
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(png_path), width=Cm(width_cm))
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_after = Pt(2)
        self._caption("figure", self.n_fig, caption, above=False)
        return self.n_fig

    # ------------------------------------------------------------------ three-line table
    @staticmethod
    def _border(cell, **edges):
        tcPr = cell._tc.get_or_add_tcPr()
        b = tcPr.find(qn("w:tcBorders"))
        if b is None:
            b = OxmlElement("w:tcBorders")
            tcPr.append(b)
        for edge, (sz, val) in edges.items():
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), val); el.set(qn("w:sz"), str(sz)); el.set(qn("w:space"), "0"); el.set(qn("w:color"), "000000")
            b.append(el)

    def table(self, df: pd.DataFrame, caption: str, note: Optional[str] = None, font: Optional[float] = None, col_widths: Optional[Sequence[float]] = None,
              left_cols: int = 1) -> int:
        self.n_tab += 1
        font = font or self.table_pt
        self._caption("table", self.n_tab, caption, above=True)
        tb = self.doc.add_table(rows=1, cols=len(df.columns))
        tb.alignment = WD_TABLE_ALIGNMENT.CENTER
        tb.autofit = False
        for j, c in enumerate(df.columns):
            cell = tb.rows[0].cells[j]
            cell.text = ""
            self._font(cell.paragraphs[0].add_run(str(c)), font, True)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if j < left_cols else WD_ALIGN_PARAGRAPH.CENTER
            self._border(cell, top=(12, "single"), bottom=(6, "single"))
        last = len(df) - 1
        for i, row in enumerate(df.itertuples(index=False)):
            cells = tb.add_row().cells
            for j, v in enumerate(row):
                cells[j].text = ""
                txt = "" if (v is None or (isinstance(v, float) and np.isnan(v))) else str(v)
                self._font(cells[j].paragraphs[0].add_run(txt), font)
                cells[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if j < left_cols else WD_ALIGN_PARAGRAPH.CENTER
                if i == last:
                    self._border(cells[j], bottom=(12, "single"))
        for r in tb.rows:
            for j, cell in enumerate(r.cells):
                cell.paragraphs[0].paragraph_format.space_before = Pt(1)
                cell.paragraphs[0].paragraph_format.space_after = Pt(1)
                if col_widths:
                    cell.width = Cm(col_widths[j])
        self._p(note or "", 9 if self.profile == "thesis" else 8, color="444444", after=8 if note else 4)
        return self.n_tab

    def save(self, path: str | Path) -> Path:
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(path)
        return path
