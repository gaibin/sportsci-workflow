"""
Tests for sportsci_figures.pubstyle — the two-language publication style.

They render the complete demo (6 figures, 3 tables, docx, xlsx) in BOTH language sets and check the
invariants an AI agent must never break:
  * every expected file exists (PNG 300 dpi + vector PDF per figure, docx, xlsx)
  * no missing-glyph warnings (no "tofu" squares) in either set
  * English deliverables contain zero CJK characters (figures, tables, captions); Chinese deliverables do contain CJK
  * captions follow the language ("Figure 1." / "图 1", "Table 1." / "表 1")
  * tables are three-line: thick top/bottom rules, thin header rule, no vertical borders
  * fonts are embedded in the PDF; English PDFs do not embed a CJK font
  * colour roles are stable
"""
from __future__ import annotations

import importlib.util
import re
import warnings
from pathlib import Path

import matplotlib.pyplot as plt
import pytest
from docx import Document
from docx.oxml.ns import qn

from sportsci_figures.pubstyle import apply_style, condition_colors, dose_colors, has_cjk_font, save_figure, verify_language
from sportsci_figures.pubstyle.style import C_AMBER, C_MAGENTA, font_chain

ROOT = Path(__file__).resolve().parents[1]
CJK = re.compile(r"[㐀-鿿　-〿＀-￯]")
FIGS = ["fig1_flow", "fig2_timecourse", "fig3_paired_estimate", "fig4_prepost", "fig5_forest", "fig6_composition_profile_scatter"]

needs_cjk = pytest.mark.skipif(not has_cjk_font(), reason="no CJK font installed; the zh set cannot be rendered on this machine")


@pytest.fixture(scope="module")
def demo(tmp_path_factory):
    out = tmp_path_factory.mktemp("pubstyle_demo")
    spec = importlib.util.spec_from_file_location("pubstyle_demo", ROOT / "examples" / "pubstyle_demo.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    msgs = {}
    for lang in (["zh", "en"] if has_cjk_font() else ["en"]):
        with warnings.catch_warnings(record=True) as w:
            warnings.simplefilter("always")
            mod.run(lang, out / lang)
        msgs[lang] = [str(x.message) for x in w]
    return out, msgs


def _doc_text(path):
    d = Document(path)
    txt = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            txt += [c.text for c in r.cells]
    return d, "\n".join(txt)


@pytest.mark.parametrize("lang", ["en", "zh"])
def test_outputs_exist_and_png_is_300dpi(demo, lang):
    out, _ = demo
    if lang == "zh" and not has_cjk_font():
        pytest.skip("no CJK font")
    from PIL import Image
    for f in FIGS:
        png, pdf = out / lang / "figures" / f"{f}.png", out / lang / "figures" / f"{f}.pdf"
        assert png.exists() and pdf.exists() and pdf.stat().st_size > 5000
        dpi = Image.open(png).info.get("dpi", (0, 0))
        assert round(dpi[0]) == 300
    assert (out / lang / "results.docx").exists()
    assert (out / lang / "tables" / "tables.xlsx").exists()


@pytest.mark.parametrize("lang", ["en", "zh"])
def test_no_missing_glyph_warnings(demo, lang):
    _, msgs = demo
    if lang not in msgs:
        pytest.skip("no CJK font")
    bad = [m for m in msgs[lang] if "missing from font" in m or "not found" in m.lower() and "font" in m.lower()]
    assert not bad, bad[:3]


def test_english_docx_has_no_cjk_and_english_captions(demo):
    out, _ = demo
    d, text = _doc_text(out / "en" / "results.docx")
    assert not CJK.search(text)
    assert "Figure 1." in text and "Table 1." in text and "Simulated-data notice" in text


@needs_cjk
def test_chinese_docx_has_chinese_captions(demo):
    out, _ = demo
    d, text = _doc_text(out / "zh" / "results.docx")
    assert "图 1" in text and "表 1" in text and "模拟数据声明" in text
    assert len(d.tables) == 3 and len(d.inline_shapes) == 6


@pytest.mark.parametrize("lang", ["en", "zh"])
def test_three_line_tables(demo, lang):
    out, _ = demo
    if lang == "zh" and not has_cjk_font():
        pytest.skip("no CJK font")
    d = Document(out / lang / "results.docx")
    for t in d.tables:
        hdr = t.rows[0].cells[0]._tc.tcPr.find(qn("w:tcBorders"))
        assert hdr.find(qn("w:top")).get(qn("w:sz")) == "12" and hdr.find(qn("w:bottom")).get(qn("w:sz")) == "6"
        last = t.rows[-1].cells[0]._tc.tcPr.find(qn("w:tcBorders"))
        assert last.find(qn("w:bottom")).get(qn("w:sz")) == "12"
        for r in t.rows:
            for c in r.cells:                              # no vertical rules anywhere
                b = c._tc.tcPr.find(qn("w:tcBorders"))
                if b is not None:
                    assert b.find(qn("w:left")) is None and b.find(qn("w:right")) is None


def test_pdf_fonts_embedded_and_language_specific(demo):
    fitz = pytest.importorskip("pymupdf")
    out, _ = demo
    fonts_en = {f[3] for f in fitz.open(out / "en" / "figures" / "fig3_paired_estimate.pdf")[0].get_fonts()}
    assert any("Arial" in f or "DejaVu" in f or "Liberation" in f or "Helvet" in f for f in fonts_en)
    assert not any("YaHei" in f or "CJK" in f or "SimHei" in f or "PingFang" in f for f in fonts_en)
    if has_cjk_font():
        fonts_zh = {f[3] for f in fitz.open(out / "zh" / "figures" / "fig3_paired_estimate.pdf")[0].get_fonts()}
        assert any(k in f for f in fonts_zh for k in ("YaHei", "CJK", "SimHei", "PingFang", "Heiti", "Source", "WenQuanYi"))


def test_english_figure_with_cjk_is_rejected(tmp_path):
    apply_style("en")
    fig, ax = plt.subplots()
    ax.set_xlabel("时间 (h)")
    assert verify_language(fig)
    with pytest.raises(ValueError):
        save_figure(fig, tmp_path / "bad")
    plt.close("all")


def test_colour_roles_are_stable():
    col = condition_colors(["CON", "RT20", "RT30", "RT45"])
    assert col["CON"] == C_AMBER and col["RT45"] == C_MAGENTA
    ramp = dose_colors(4)
    lum = [0.299 * int(c[1:3], 16) + 0.587 * int(c[3:5], 16) + 0.114 * int(c[5:7], 16) for c in ramp]
    assert lum == sorted(lum, reverse=True)           # lightness decreases with dose: survives greyscale printing


def test_font_chain_en_has_no_cjk_and_zh_falls_back():
    assert not any(f in ("Microsoft YaHei", "SimHei", "PingFang SC", "Noto Sans CJK SC") for f in font_chain("en"))
    if has_cjk_font():
        chain = font_chain("zh")
        assert chain[0] in ("Arial", "Helvetica", "Liberation Sans")     # Latin first => Latin text is Arial
        assert any(f not in ("Arial", "Helvetica", "Liberation Sans") for f in chain[1:])
